from pathlib import Path
import re, json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parent
# Authoritative content is the public October 21 CV source, preserved in content.json.
original_records=json.loads((ROOT/'content.json').read_text())
sections={}; front=[]; section=None
for r in original_records:
    if r['type']=='heading': section=r['text']; sections[section]=[]
    if section: sections[section].append(r.copy())
    else: front.append(r.copy())
ORDER=['Research Profile','Research Interests','Education','Publications','Conference Presentations','Awards & Honors','Academic Service','Research Projects','Teaching Experience','Technical Skills']
records=front[:]
def rec(typ,text,**kw): return dict(type=typ,text=text,**kw)
for name in ORDER:
    old=sections[name]; block=[old[0]]
    if name=='Education':
        institution=location=None
        for r in old[1:]:
            if r['type']=='pair': institution,location=r['text'].split('\t')
            elif re.match(r'^(Ph.D.|B.S.|B.Eng.)',r['text']):
                degree,date=re.split(r'\s{2,}',r['text'])
                block.extend([rec('entry',degree),rec('metadata',institution+' | '+location+'\t'+date)])
            else: block.append(r)
    elif name=='Teaching Experience':
        for i in range(1,len(old),2):
            course,institution=old[i]['text'].split('\t')
            role,date=old[i+1]['text'].split('\t')
            block.extend([rec('entry',course),rec('metadata',institution+'\t'+date),rec('role',role)])
    elif name=='Awards & Honors':
        title,institution=old[1]['text'].split('\t'); role,date=old[2]['text'].split('\t')
        block.extend([rec('entry',title),rec('metadata',institution+'\t'+date),rec('role',role)])
        block.extend(old[3:])
    elif name=='Academic Service':
        for r in old[1:]:
            left,date=r['text'].rsplit(' | ',1)
            block.append(rec('service',left+'\t'+date))
    else:
        for r in old[1:]:
            if r['type']=='project': r['type']='entry'
            elif r['type']=='pair': r['type']='metadata'
            block.append(r)
    records.extend(block)
(ROOT/'records.json').write_text(json.dumps(records,indent=2,ensure_ascii=False))

D=Document(); sec=D.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=Inches(.52); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(1); sec.right_margin=Inches(1)
sec.header_distance=Inches(.2); sec.footer_distance=Inches(.3)
for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
    sty=D.styles[name]; sty.font.name='Times New Roman'; sty.font.color.rgb=RGBColor(0,0,0)
    sty._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Times New Roman')
normal=D.styles['Normal']; normal.font.size=Pt(10.5)
normal.paragraph_format.line_spacing=Pt(15.6)
normal.paragraph_format.space_after=Pt(0)
normal.paragraph_format.widow_control=True
D.core_properties.title='Curriculum Vitae/Resume'
D.core_properties.author='Liuyang Cheng'
D.core_properties.subject='Curriculum vitae'
D.core_properties.keywords=''

# Original document highlights the CV owner's name and journal/conference titles.
italics=['Advanced Intelligent Systems','Advanced Science','International Journal of Mechanical Sciences','Proceedings of the National Academy of Sciences','Smart Materials, Adaptive Structures, and Intelligent Systems','Society of Engineering Science','Advanced Robotics Research']

def hyperlink(p,text,url):
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r'); rp=OxmlElement('w:rPr')
    f=OxmlElement('w:rFonts'); f.set(qn('w:ascii'),'Times New Roman'); f.set(qn('w:hAnsi'),'Times New Roman'); rp.append(f)
    color=OxmlElement('w:color'); color.set(qn('w:val'),'16A2D6'); rp.append(color)
    u=OxmlElement('w:u'); u.set(qn('w:val'),'single'); rp.append(u)
    r.append(rp); t=OxmlElement('w:t'); t.text=text; r.append(t); h.append(r); p._p.append(h)

def text_runs(p,text,all_bold=False,all_italic=False,size=None):
    regex='('+ '|'.join(re.escape(s) for s in ['https://doi.org/10.1002/aisy.70490','Cheng, L.','Liuyang Cheng']+italics) +')'
    for s in re.split(regex,text):
        if not s: continue
        if s.startswith('https://'): hyperlink(p,s,s); continue
        r=p.add_run(s); r.font.name='Times New Roman'; r.bold=all_bold or s in ('Cheng, L.', 'Liuyang Cheng') and text!=s; r.italic=all_italic or s in italics
        if size: r.font.size=Pt(size)

def border(p):
    ppr=p._p.get_or_add_pPr(); b=OxmlElement('w:pBdr'); bottom=OxmlElement('w:bottom')
    for k,v in [('val','single'),('sz','4'),('space','2'),('color','000000')]: bottom.set(qn('w:'+k),v)
    b.append(bottom); ppr.append(b)

def bullet_setup(p):
    f=p.paragraph_format; f.left_indent=Inches(.5); f.first_line_indent=Inches(-.25)
    f.tab_stops.add_tab_stop(Inches(.5))
    r=p.add_run('○\t'); r.font.name='Times New Roman'; r.font.size=Pt(8)

def pair(p,text,section):
    left,right=text.split('\t')
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5),WD_TAB_ALIGNMENT.RIGHT)
    text_runs(p,left,size=10)
    p.add_run('\t'); text_runs(p,right,all_italic=True,size=10)

section=''; prev=None
for ix,rec in enumerate(records):
    typ,t=rec['type'],rec['text']
    p=D.add_paragraph(); pf=p.paragraph_format
    if typ=='title':
        p.style='Title'; pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.space_after=Pt(1); pf.line_spacing=1; text_runs(p,t,all_bold=True,size=16); pf.keep_with_next=True
    elif typ=='name':
        pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.line_spacing=1; pf.space_after=Pt(0); text_runs(p,t,size=11); pf.keep_with_next=True
    elif typ=='email':
        pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; hyperlink(p,t,'mailto:'+t); pf.space_after=Pt(18); pf.keep_with_next=True
    elif typ=='heading':
        section=t; p.style='Heading 1'; pf.space_before=Pt(14 if t!='Research Profile' else 0); pf.space_after=Pt(3); pf.line_spacing=1; pf.keep_with_next=True; text_runs(p,t,all_bold=True,size=12); border(p)
    elif typ=='subheading':
        pf.left_indent=Inches(.25); pf.space_before=Pt(5); pf.space_after=Pt(3); pf.keep_with_next=True; text_runs(p,t,all_bold=True,all_italic=True)
    elif typ=='body':
        p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; text_runs(p,t)
    elif typ=='bullet':
        bullet_setup(p); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; pf.space_after=Pt(0); pf.keep_together=True
        if section=='Technical Skills' and ':' in t:
            lead,rest=t.split(':',1); text_runs(p,lead+':',all_bold=True); text_runs(p,rest)
        else: text_runs(p,t,size=10 if section in ('Publications','Research Projects') else None)
        if section=='Conference Presentations': pf.space_after=Pt(3)
    elif typ=='entry':
        pf.space_before=Pt(11 if prev and prev['type']!='heading' else 0); pf.keep_with_next=True; text_runs(p,t,all_bold=True)
    elif typ=='metadata':
        pair(p,t,section); pf.keep_with_next=True
    elif typ=='service':
        pair(p,t,section); pf.keep_together=True
    elif typ=='role':
        text_runs(p,t,all_italic=True); pf.keep_with_next=section!='Teaching Experience'
    else:
        text_runs(p,t); pf.keep_with_next=True
    prev=rec

# Keep each research entry intact, rather than letting a project straddle a page.
active=False
for i,r in enumerate(records):
    if r['type']=='heading': active=r['text']=='Research Projects'
    if active and i+1<len(records) and records[i+1]['type'] not in ('entry','heading'):
        D.paragraphs[i].paragraph_format.keep_with_next=True

# Stable section-boundary pagination in a four-page CV.
for p in D.paragraphs:
    if p.text in ['Accepted','Teaching Experience']:
        p.paragraph_format.page_break_before=True

# Remove template title rules and theme fonts so the source's Times style survives rendering.
for sty in D.styles:
    for tag in ['w:asciiTheme','w:hAnsiTheme','w:cstheme','w:eastAsiaTheme']:
        for el in sty._element.iter(qn('w:rFonts')):
            el.attrib.pop(qn(tag),None)
for node in [D.styles['Title']._element, D.paragraphs[0]._p]:
    for border_node in list(node.iter(qn('w:pBdr'))):
        border_node.getparent().remove(border_node)
out=ROOT/'Liuyang-Cheng-CV-en-20261022.docx'
D.save(out)
print(out)
for i,r in enumerate(records): print(i,r['type'],repr(r['text']))
