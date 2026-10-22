from pathlib import Path
from copy import deepcopy
import argparse, re, json
from lxml import html, etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parent
ap=argparse.ArgumentParser();ap.add_argument('--source',type=Path,default=ROOT/'cv-zh.html');ap.add_argument('--output-dir',type=Path,default=ROOT/'output');args=ap.parse_args()
args.output_dir.mkdir(parents=True,exist_ok=True)
tree=html.parse(str(args.source)); root=tree.getroot(); main=root.find('.//main')
# Flatten old printed pages, omit running furniture and redundant continuation heading.
blocks=[]
for sec in (main.findall('section') or [main]):
    for el in sec:
        if el.tag=='footer' or 'running-header' in el.get('class','') or (el.tag=='h2' and el.text_content()=='科研项目（续）'):continue
        blocks.append(deepcopy(el))
# Reorder complete sections, preserving every other section's relative order.
header=blocks[0]; sections=[]
for el in blocks[1:]:
    if el.tag=='h2':sections.append([el])
    else:sections[-1].append(el)
by={s[0].text_content():s for s in sections}
sections=[s for s in sections if s[0].text_content() not in ('奖励与荣誉','技术技能')]
i=next(i for i,s in enumerate(sections) if s[0].text_content()=='学术服务');sections.insert(i,by['奖励与荣誉']);sections.append(by['技术技能'])
for c in list(main): main.remove(c)
main.append(header)
for s in sections:
    for e in s:main.append(e)
# Canonical source is kept as faithful content; builder applies the unified metadata layout.
css=root.find('.//style');css.text='''@page {size:Letter; margin:36pt 43.2pt;} body {font-family:"Noto Sans CJK SC",sans-serif;font-size:9.5pt;line-height:1.28;color:#000} header {text-align:center} h1{font-size:21pt} h2{font-size:12pt;margin:10pt 0 4pt} h3{font-size:10pt;margin:6pt 0 1pt} p{margin:0 0 3pt} .row{display:flex;justify-content:space-between;gap:12pt} .date{white-space:nowrap} ul{padding-left:14pt;margin:2pt 0 5pt} .en-title{font-size:8.4pt} .project,.education,.award{break-inside:avoid} a{color:inherit;text-decoration:none}'''
args.output_dir.joinpath('cv-zh.html').write_bytes(etree.tostring(root,encoding='utf-8',method='html',doctype='<!doctype html>'))

doc=Document(); sec=doc.sections[0]; sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=Inches(.48);sec.bottom_margin=Inches(.48);sec.left_margin=sec.right_margin=Inches(.6);sec.footer_distance=Inches(.2)
FONT='Noto Sans CJK SC'; LATIN='Noto Sans'
def setfont(run,size=None,bold=None,italic=None):
    run.font.name=LATIN;run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT)
    if size:run.font.size=Pt(size)
    if bold is not None:run.bold=bold
    if italic is not None:run.italic=italic
    run.font.color.rgb=RGBColor(0,0,0)
for name,size,bold in [('Normal',9.4,False),('Title',21,True),('Heading 1',12,True),('Heading 2',10,True)]:
    st=doc.styles[name];st.font.name=LATIN;st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT);st.font.size=Pt(size);st.font.bold=bold;st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_after=Pt(3);st.paragraph_format.line_spacing=Pt(size*1.35)
normal=doc.styles['Normal'];normal.paragraph_format.widow_control=True
for name in ['Heading 1','Heading 2']:
    doc.styles[name].paragraph_format.keep_with_next=True
    doc.styles[name].paragraph_format.space_before=Pt(9 if name=='Heading 1' else 5)
    doc.styles[name].paragraph_format.space_after=Pt(3 if name=='Heading 1' else 1)

def text(el):return re.sub(r'\s+',' ',el.text_content()).strip()
def pnew(style=None,size=9.4,before=0,after=3,keep=False):
    p=doc.add_paragraph(style=style);p.paragraph_format.space_before=Pt(before);p.paragraph_format.space_after=Pt(after);p.paragraph_format.keep_with_next=keep;p.paragraph_format.line_spacing=Pt(size*1.35)
    return p

def addplain(p,t,size=9.4,bold=False,italic=False):
    r=p.add_run(t);setfont(r,size,bold,italic);return r

def inline(p,node,size=9.4,bold=False,italic=False):
    # Inline semantic formatting and explicit hyperlinks remain editable.
    b=bold or node.tag in ('strong','b');i=italic or node.tag in ('em','i')
    if node.text:addplain(p,node.text,size,b,i)
    for c in node:
        if c.tag=='br':p.add_run().add_break()
        elif c.tag=='a' and c.get('href'):
            h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(c.get('href'),RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');pr=OxmlElement('w:rPr');f=OxmlElement('w:rFonts');f.set(qn('w:ascii'),LATIN);f.set(qn('w:hAnsi'),LATIN);f.set(qn('w:eastAsia'),FONT);pr.append(f);sz=OxmlElement('w:sz');sz.set(qn('w:val'),str(int(size*2)));pr.append(sz);r.append(pr);t=OxmlElement('w:t');t.text=c.text_content();r.append(t);h.append(r);p._p.append(h)
        else:inline(p,c,size,b,i)
        if c.tail:addplain(p,c.tail,size,b,i)

def paragraph(el,size=9.4,keep=False,after=3):
    p=pnew(size=size,keep=keep,after=after);inline(p,el,size);return p

def title(t,size=10):
    p=pnew(size=size,keep=True,before=5,after=1);addplain(p,t,size,True);return p

def metadata(left,right,after=2):
    p=pnew(size=9,keep=True,after=after);p.paragraph_format.tab_stops.add_tab_stop(Inches(7.3),WD_TAB_ALIGNMENT.RIGHT)
    addplain(p,left,9);addplain(p,'\t'+right,9);return p

def bullet(el,size=9.4,after=2):
    p=pnew(size=size,after=after);p.paragraph_format.left_indent=Pt(12);p.paragraph_format.first_line_indent=Pt(-9);addplain(p,'•\t',size);p.paragraph_format.tab_stops.add_tab_stop(Pt(12));inline(p,el,size);return p

# Header, using exact user-supplied name and contact.
for e in header:
    p=pnew(style='Title' if e.tag=='h1' else None,size=21 if e.tag=='h1' else 10,after=3);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;inline(p,e,21 if e.tag=='h1' else 10 if e.tag=='div' else 9.4,bold=e.tag=='h1')

for group in sections:
    heading=group[0].text_content(); p=doc.add_paragraph(heading,'Heading 1')
    if heading=='会议报告':p.paragraph_format.page_break_before=True
    for el in group[1:]:
        cl=el.get('class','').split()
        if 'education' in cl:
            schoolrow=el.find('div');school=text(schoolrow[0]);loc=text(schoolrow[1]);items=el.find('ul').findall('li')
            for item in items:
                row=item.find('div')
                if row is not None:
                    title(text(row[0]));metadata(school+' · '+loc,text(row[1]))
                else:paragraph(item,9.2,after=2)
        elif 'bibliography' in cl or (el.tag=='ul' and heading=='会议报告'):
            for c in el:
                if c.tag=='h3':title(text(c),9.8)
                elif c.tag=='ul':
                    for li in c:bullet(li,8.6,after=3)
                elif c.tag=='li':bullet(c,8.6,after=4)
        elif 'project' in cl:
            for c in el:
                ccl=c.get('class','').split()
                if c.tag=='h3':
                    p=title(text(c))
                    if text(c)=='拉伸 J 型曲线材料在拮抗构型中的应用原理':p.paragraph_format.page_break_before=True
                elif 'en-title' in ccl:paragraph(c,8.4,keep=True,after=1)
                elif 'row' in ccl:metadata(text(c[0]),text(c[1]))
                elif c.tag=='ul':
                    for li in c:bullet(li,9.2,after=1.8)
                else:paragraph(c,9.1,keep=True,after=1.5)
        elif 'teaching' in cl:
            for row in el:
                course=row.xpath('.//*[contains(@class,"course")]')[0];school=row.xpath('.//*[contains(@class,"school")]')[0];date=row.xpath('.//*[contains(@class,"date")]')[0]
                title(text(course)+' · 助教',9.6);metadata(text(school),text(date),after=3)
        elif 'award' in cl:
            title(text(el[0][0]),9.8);metadata(text(el[1]),text(el[0][1]));
            for c in el[2:]:paragraph(c,9.2,after=2)
        elif el.tag=='ul':
            for li in el:bullet(li,9.2,after=2.5)
        elif heading=='学术服务':
            a,b=text(el).rsplit('|',1);metadata(a.strip(),b.strip(),after=2)
        else:paragraph(el)

# Running footer, dynamic page fields.
p=sec.footer.paragraphs[0];p.paragraph_format.tab_stops.clear_all();p.paragraph_format.tab_stops.add_tab_stop(Inches(7.3),WD_TAB_ALIGNMENT.RIGHT);addplain(p,'Liuyang Cheng · 学术简历\t',8)
for label,code in [('', 'PAGE'),(' / ','NUMPAGES')]:
    addplain(p,label,8);r=p.add_run();setfont(r,8);fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),code);r._r.addnext(fld)
# Remove template title rules and Asian grid snapping.
for e in list(doc.styles.element.iter()):
    if e.tag in (qn('w:pBdr'),qn('w:shd')):e.getparent().remove(e)
for e in list(sec._sectPr):
    if e.tag==qn('w:docGrid'):sec._sectPr.remove(e)
for p in doc.paragraphs:
    e=OxmlElement('w:snapToGrid');e.set(qn('w:val'),'0');p._p.get_or_add_pPr().append(e)
out=args.output_dir/'Liuyang-Cheng-CV-zh-20261022.docx';doc.save(out)
print(out)
print('headings', [s[0].text_content() for s in sections])
print('paragraphs',len(doc.paragraphs))
